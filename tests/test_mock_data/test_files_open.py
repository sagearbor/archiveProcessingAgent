from pathlib import Path
import zipfile
import tarfile
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

DATA_DIR = Path(__file__).resolve().parent.parent.parent / "mock_data"


def test_open_word():
    Document(DATA_DIR / "mock_word.docx")


def test_open_excel():
    load_workbook(DATA_DIR / "mock_excel.xlsx")


def test_open_pptx():
    Presentation(DATA_DIR / "mock_powerpoint.pptx")


def test_open_powerbi():
    with zipfile.ZipFile(DATA_DIR / "mock_powerbi.pbix") as z:
        assert len(z.namelist()) > 0


def test_open_tableau():
    with zipfile.ZipFile(DATA_DIR / "mock_tableau.twbx") as z:
        assert len(z.namelist()) > 0


def test_open_synapse():
    with zipfile.ZipFile(DATA_DIR / "mock_synapse.zip") as z:
        assert len(z.namelist()) > 0


def test_open_archive():
    with zipfile.ZipFile(DATA_DIR / "mock_archive.zip") as z:
        assert len(z.namelist()) > 0


def test_open_source():
    with tarfile.open(DATA_DIR / "mock_source.tar.gz") as t:
        assert len(t.getmembers()) > 0